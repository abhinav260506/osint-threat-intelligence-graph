import docx

doc_path = r"C:\Users\abhin\.gemini\antigravity\brain\c50bbb3a-7936-4b47-a876-0b82d38cf043\.user_uploaded\media_1793375304024_07c11224.docx"
doc = docx.Document(doc_path)

def check_text(p_idx, text, source_name='p'):
    terms = ['louvain', 'zero-hallucination', 'cryptographic provenance', 'sub-millisecond', '7.47s', '81 nodes']
    for t in terms:
        if t in text.lower():
            print(f'[{source_name} {p_idx}] contains "{t}":\n  "{text.strip()[:140]}"')

for i, p in enumerate(doc.paragraphs):
    check_text(i, p.text, 'Para')

for i, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            check_text(f'T{i} R{r_idx}C{c_idx}', cell.text, 'Cell')

